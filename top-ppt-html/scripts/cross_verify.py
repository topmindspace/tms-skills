#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""TopPPT HTML · 双通道交叉验证（可选依赖 python-pptx）：
① A 通道（手写 OOXML）与 B 通道（pptxgenjs）产物均可被 python-pptx 严格解析
② 两通道逐页文本集合一致（同一模型 → 同一内容；段落换行表示差异已归一）
   归一化：B 通道带数据的图表为原生 chart part（可编辑数据）——图表类别标签并入
   文本集合、纯数值 token（含百分数）两端过滤（A 通道形状版的数值标签 vs B 通道 chart 数据值）
③ 文本溢出启发式：文本估算宽 vs shape 宽 × 可容行数
④ B 通道原生图表数值 ↔ 模型数值（类别标签一致不代表数值一致；数值错乱只能这里拓）
⑤ 字号比例尺同源：B 的字号必须是 A 字号集的子集（防 modeSize()/sz() 后处理漂移）
默认 SKIP；加 --full-ab 才执行。用法：先 node scripts/gen_channel_a.js 生成 A 通道产物，再 python scripts/cross_verify.py
改 pptx-export.js 序列化骨架后必跑（见 references/pptx-export.md 双裁判教训）。
"""
import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
# 依赖缺失不在导入期退出——本模块的核对函数会被 negative_tests.py 直接复用
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

ROOT = Path(__file__).resolve().parent.parent
A_DIR = ROOT / 'dist' / 'regression-a'    # A 通道（页面直导引擎）
B_DIR = ROOT / 'dist' / 'regression'      # B 通道（pptxgenjs）
MODEL_DIR = ROOT / 'assets' / 'examples'

# 文本度量单源（与 build_pptx.js / validate_pptx.py 同值）——三侧各自硬编码会让
# 溢出判定在不同裁判间飘移，出现“一侧报溢出、一侧说没事”。
try:
    _TM = (json.loads((ROOT / 'scripts' / 'layout-constants.json').read_text(encoding='utf-8'))
           .get('containers', {}).get('textMetrics', {}))
except (OSError, json.JSONDecodeError):
    _TM = {}
LINE_FACTOR = float(_TM.get('lineFactor') or 1.45)
EM_ASCII_RATIO = float(_TM.get('emAsciiRatio') or 0.52)

# 形态变换类图表：模型里的负值在图上是「向下的正高度条」，按绝对值比对
SIGNLESS_CHARTS = {'waterfall', 'gauge', 'funnel', 'bullet'}

# 纯数值 token（含千分位/小数/百分号/常见单位后缀）——图表数据标签在 A 通道是
# 文本（值+unit，如 "9pt"），B 通道进 chart part；比对前两端过滤。类别标签（如 "2026Q3"）
# 不匹配此模式，仍参与一致性检查。
NUMERIC_TOKEN = re.compile(
    r'^[+-]?[\d.,]+\s*(?:%|‰|pt|pct|pp|YB|ZB|EB|PB|TB|GB|MB|KB|kWh|Gbps|k|w|万|亿|元|倍|个|人|天|家|次|项|年'
    r'|条|笔|台|套|份|款|张|篇|页|件|例|场|轮|艘|辆|吨|米|秒|分|时|档|级)?$')


def slide_texts(prs):
    """每页文本集合（逐行拆分——A 通道段落 \\n 与 B 通道 breakLine \\n\\n 表示差异归一；
    忽略空白、页码与纯数值 token；B 通道原生图表的类别标签并入集合）。"""
    pages = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for line in shape.text_frame.text.splitlines():
                    t = line.strip()
                    if t and not NUMERIC_TOKEN.match(t):
                        texts.append(t)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for line in cell.text.splitlines():
                            t = line.strip()
                            if t and not NUMERIC_TOKEN.match(t):
                                texts.append(t)
            elif getattr(shape, 'has_chart', False):
                # B 通道原生数据图表——类别标签（实质内容）并入；数据值已按纯数值过滤。
                # 类别标签同样过 NUMERIC_TOKEN——A 通道文本统一过滤，两端语义对称
                # （否则 gauge 的 '76%'、stack 的年份 '2023' 这类数值型类别会造成单边误报）。
                try:
                    cats = list(shape.chart.plots[0].categories)
                    texts.extend(str(c).strip() for c in cats
                                 if str(c).strip() and not NUMERIC_TOKEN.match(str(c).strip()))
                except Exception:
                    pass
        pages.append(sorted(set(texts)))   # 集合语义去重（B 端 chart categories 与自绘图例可能重复）
    return pages


def _model_chart(sec):
    """取章节页的图表对象（口径与 validate_pptx._section_chart 一致）。"""
    st = sec.get('type') or ''
    if st in {'bar', 'donut', 'exhibit', 'halftable'}:
        c = sec.get('chart') or {}
    elif st == 'split':
        c = sec.get('right') or {}
        if (c.get('type') or 'bar') == 'table':
            return None
    else:
        return None
    return c if isinstance(c, dict) and c.get('labels') and c.get('values') else None


def _model_values(chart):
    """模型声明的全部数值（单系列 values 或多系列 series[].values）。"""
    out = []
    if isinstance(chart.get('series'), list) and chart['series']:
        for se in chart['series']:
            out += [v for v in (se.get('values') or []) if isinstance(v, (int, float))]
    else:
        out += [v for v in (chart.get('values') or []) if isinstance(v, (int, float))]
    return out


def chart_data_verify(prs, model, name):
    """B 通道原生 chart part 的数值必须覆盖模型声明值。

    只做「模型值 ⊆ 图表值」的包含判定，不做严格相等——waterfall 的累计基座、
    gauge 的补角这类派生系列是合法的额外数据，严格相等会制造假阳性。
    形态变换类图表（waterfall 把 -3 画成高度 3 的悬浮条）按绝对值比对。
    """
    issues = []
    secs = [s for s in (model.get('sections') or []) if isinstance(s, dict)]
    first = 3 if model.get('agenda') else 2
    slides = list(prs.slides)
    for i, sec in enumerate(secs):
        chart = _model_chart(sec)
        if chart is None:
            continue
        idx = first + i - 1
        if idx >= len(slides):
            break
        pool = []
        for shape in slides[idx].shapes:
            if not getattr(shape, 'has_chart', False):
                continue
            try:
                for plot in shape.chart.plots:
                    for se in plot.series:
                        pool += [v for v in se.values if v is not None]
            except Exception:
                pass
        if not pool:
            continue          # 该页是形状还原通道，数值走数据表/备注，另有门禁
        ctype = str(chart.get('type') or 'bar').lower()
        signless = ctype in SIGNLESS_CHARTS
        if signless:
            pool = [abs(p) for p in pool]
        missing = []
        for v in _model_values(chart):
            target = abs(v) if signless else v
            if not any(abs(target - p) <= max(0.01, abs(target) * 0.001) for p in pool):
                missing.append(v)
        if missing:
            issues.append(f'{name} 第{idx + 1}页 {ctype} 图表数值与模型不符，'
                          f'缺 {missing[:5]}（图表实际含 {sorted(set(pool))[:6]}）')
    return issues


def font_sizes(prs):
    """全文文本框出现过的字号集合（pt）。"""
    out = set()
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.text.strip():
                        out.add(round(r.font.size.pt, 2))
    return out


def overflow_estimate(prs, name):
    """启发式：文本估算宽（CJK 字符按 1em、ASCII 按 emAsciiRatio）vs shape 宽 × 可容行数。"""
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.width or not shape.height:
                continue
            for para in shape.text_frame.paragraphs:
                line = ''.join(r.text for r in para.runs)
                if not line.strip():
                    continue
                sz = None
                for r in para.runs:
                    if r.font.size:
                        sz = r.font.size.pt
                        break
                sz = sz or 12
                emu_w = int(shape.width)
                in_w = emu_w / 914400
                # 每行可容字符宽（英寸）→ em 数
                est_w = sum(1.0 if ord(c) > 0x2E80 else EM_ASCII_RATIO for c in line) * sz / 72
                lines_avail = max(1, int(int(shape.height) / 914400 / (sz / 72 * LINE_FACTOR)))
                if est_w > in_w * lines_avail * 1.18:   # 18% 容差
                    issues.append(f'{name} 第{si}页 "{line[:14]}…" 估算{est_w:.1f}in > 容量{in_w*lines_avail:.1f}in')
    return issues


def main() -> int:
    if Presentation is None:
        print('python-pptx 未安装')
        return 2
    fails = []
    a_files = sorted(A_DIR.glob('*.pptx'))
    print(f'A 通道产物 {len(a_files)} 个 · B 通道产物 {len(sorted(B_DIR.glob("*.pptx")))} 个\n')
    all_overflow = []

    for ap in a_files:
        bp = B_DIR / ap.name
        name = ap.stem
        # ① 可解析性
        try:
            prs_a = Presentation(str(ap))
        except Exception as e:
            fails.append(f'{name}: A 通道 python-pptx 解析失败 {e}')
            continue
        if not bp.exists():
            fails.append(f'{name}: B 通道产物缺失')
            continue
        try:
            prs_b = Presentation(str(bp))
        except Exception as e:
            fails.append(f'{name}: B 通道 python-pptx 解析失败 {e}')
            continue
        # ② 逐页文本一致性（v9：A/B 不必逐字全等——A 预览为形状近似；
        #    要求「标题+主文本 ≥80% 交集覆盖」，杜绝截断/串页，又不因图表标签形态差误杀）
        ta, tb = slide_texts(prs_a), slide_texts(prs_b)
        if len(ta) != len(tb):
            fails.append(f'{name}: 页数不一致 A={len(ta)} B={len(tb)}')
            continue
        min_cov = 0.80
        try:
            min_cov = float((json.loads((ROOT / 'scripts' / 'layout-constants.json')
                                       .read_text(encoding='utf-8'))
                             .get('qualityGates') or {}).get('roundtripMin') or 0.8)
        except Exception:
            pass
        diff_pages = []
        for i, (x, y) in enumerate(zip(ta, tb), 1):
            sx, sy = set(x), set(y)
            if not sx and not sy:
                continue
            inter = sx & sy
            cov = len(inter) / max(1, max(len(sx), len(sy)))
            if cov < min_cov:
                diff_pages.append((i, cov, sorted(sx - sy)[:3], sorted(sy - sx)[:3]))
        if diff_pages:
            for i, cov, only_a, only_b in diff_pages[:3]:
                fails.append(f'{name}: 第{i}页文本覆盖 {cov:.0%} <{min_cov:.0%} '
                             f'A独有{only_a} B独有{only_b}')
        # ③ 溢出启发式（对两通道都跑）
        all_overflow += overflow_estimate(prs_a, name + '/A')
        all_overflow += overflow_estimate(prs_b, name + '/B')
        # ④ B 通道原生图表数值 ↔ 模型（A 通道是形状还原，不产 chart part）
        chart_note = ''
        mp = MODEL_DIR / f'{name}.model.json'
        if mp.exists():
            try:
                model = json.loads(mp.read_text(encoding='utf-8'))
            except json.JSONDecodeError as e:
                fails.append(f'{name}: 模型不可解析 {e}')
                model = None
            if model:
                data_issues = chart_data_verify(prs_b, model, name)
                fails += data_issues
                chart_note = ' 图表数值=' + ('PASS' if not data_issues else f'DIFF({len(data_issues)})')
        status = 'PASS' if not diff_pages else f'LOWCOV p{[d[0] for d in diff_pages]}'
        # ⑤ 字号比例尺同源：B 的字号必须都来自 A 同一张表。
        #    反向不成立——A 把图表渲染成形状，标签字号会多出来（B 的在 chart part 内）。
        fa, fb = font_sizes(prs_a), font_sizes(prs_b)
        alien = sorted(fb - fa)
        if alien:
            fails.append(f'{name}: B 通道出现 A 没有的字号 {alien}'
                         f'（两通道应同走 modeTypeScale；modeSize() 与 sz() 的钳制/取整须一致）')
        sz_note = ' 字号比例尺=' + ('PASS' if not alien else 'DRIFT')
        print(f'{name}: A可解析✓ B可解析✓ 页数={len(ta)} 逐页文本={status}{chart_note}{sz_note}')

    print('\n── 溢出启发式（check-overflow · 18% 容差）──')
    if all_overflow:
        for i in all_overflow[:20]:
            print('  [WARN]', i)
    else:
        print('  两通道全部通过（无文本超容 shape）')

    print('\n' + '=' * 50)
    if fails:
        print('交叉验证失败:')
        for f in fails:
            print('  [FAIL]', f)
        return 1
    print('双通道交叉验证通过：python-pptx 第三方裁判确认 A/B 产物均可解析、'
          '逐页文本一致、原生图表数值与模型一致')
    return 0


if __name__ == '__main__':
    # Batch 2: A/B 全文对等默认关闭；显式 --full-ab 才跑（Fast/标准交付不依赖）
    if '--full-ab' not in sys.argv:
        print('cross_verify: SKIP（默认关闭；需要 A/B 全文对等时加 --full-ab）')
        sys.exit(0)
    sys.exit(main())
